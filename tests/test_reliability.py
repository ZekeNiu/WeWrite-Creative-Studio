"""User-visible reliability contracts; only synthetic data and model replies."""
import copy
import asyncio
import io
import json
import sqlite3
import zipfile
import pytest
from docx import Document
from backend import store, creative, rendering, materials, providers, security, workflow, research, evidence_state
from backend.models import Settings, EvidenceSpan
from tests.test_studio import client, new, patch, H, wait
from tests.test_creative_flow import seeded


def span(cid='C1', sid='S1', quote='Original evidence 1'):
    return dict(claim_id=cid,source_id=sid,quote=quote,claim='主张'+cid[1:],quality='suitable',
                source_type='original',adoption_reason='Direct evidence',use_scope='Studied group',verification='quote_matched',boundary='')


def test_restore_removes_future_fields_and_rejects_active_job(client):
    a = new(client, 'Original')
    a = store.save_article(a['id'], a['revision'], lambda v: creative.adopt(v, 'Later'), 'adopt')
    version = store.versions(a['id'])[0]['id']
    restored = store.restore(a['id'], version, a['revision'])
    assert restored['brief']['topic'] == 'Original'
    assert restored.get('creative_intent', {}).get('selected', {}).get('title', 'Original') == 'Original'
    job = store.create_job(a['id'], dict(stage='write'))
    with pytest.raises(store.Conflict): store.restore(a['id'], version, restored['revision'])
    store.update_job(job['id'], status='cancelled')


def test_export_only_cited_metadata(client):
    a = new(client)
    a['sources'] = [dict(materials.source('Cited', 'SOURCE_BODY_CANARY'), id='S1'),
                    dict(materials.source('Private', 'PRIVATE_CANARY'), id='S2', selected=False, personal_material=True, use='PRIVATE_NOTE')]
    a['content'] = 'A statement [S1].'
    with zipfile.ZipFile(io.BytesIO(rendering.export_zip(a))) as z:
        data = z.read('来源清单.json').decode()
        assert all(x not in data for x in ('PRIVATE_CANARY', 'SOURCE_BODY_CANARY', 'PRIVATE_NOTE', 'personal_material'))
        rows = json.loads(data)
        assert [x['id'] for x in rows] == ['S1']


def test_settings_failed_validation_changes_nothing(client):
    cfg = Settings.model_validate(dict(services=[dict(id='a', model='m', key='old')], default_service='a'))
    providers.save_settings(cfg)
    before = providers.settings()
    broken = Settings.model_validate(dict(services=[dict(id='a', model='m', key='new'), dict(id='b', base_url='http://remote.example', model='m')], default_service='a'))
    with pytest.raises(ValueError): providers.save_settings(broken)
    assert security.key('a') == 'old'
    assert providers.settings() == before


def test_docx_keeps_document_order():
    doc = Document(); doc.add_paragraph('A before table')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'B table'
    doc.add_paragraph('C after table')
    blob = io.BytesIO(); doc.save(blob)
    text, _ = materials.extract_file('test.docx', blob.getvalue())
    assert text.index('A before') < text.index('B table') < text.index('C after')


def test_composite_unknown_citation_rejected(client):
    a = new(client); a['sources'] = [dict(materials.source('S', 'text'), id='S1')]
    with pytest.raises(ValueError): workflow.validate_result('write', 'A [S1,Sunknown]', a)


def test_unassessed_quote_is_not_support(client):
    a = new(client)
    a['research'] = dict(issues=[dict(id='Q1', text='Question', claim='Claim', claim_id='C1', kind='blocking', status='open', source_ids=['S1'])])
    span = EvidenceSpan(source_id='S1', quote='Only an association', claim='Claim', claim_id='C1').model_dump()
    notes = dict(evidence=[span], issues=[dict(a['research']['issues'][0], status='resolved', resolution='Located')])
    assert evidence_state.merge_claims(a, [span])[0]['status'] != 'supported'
    assert evidence_state.merge_issues(a, notes)[0]['status'] == 'open'


def test_article_limits_persist(client):
    a = new(client); limits = dict(max_calls=33, max_pages=40, max_rounds=3)
    a = patch(client, a, dict(research_limits=limits))
    assert client.get('/api/articles/' + a['id']).json()['research_limits'] == limits


def test_scoped_completion_and_cache_key(client):
    a=seeded(client)
    a['research']['issues'][1]['status']='open'
    j=store.create_job(a['id'],dict(stage='research',issue_ids=['Q1']))
    worker=research.Research(a,j['id'],'research')
    worker.notes=dict(evidence=[span()],issues=[dict(a['research']['issues'][0],status='resolved',resolution='Direct support')])
    assert worker.sufficient()
    assert research.input_key(a,'research','','config',['Q1'])!=research.input_key(a,'research','','config',['Q2'])


def test_current_evidence_projection_and_missing_new_support(client):
    a=seeded(client)
    a['sources'][0]['evidence_spans']=[dict(span(),quote='old')]
    a['evidence']['claims']=evidence_state.merge_claims(a,[span(quote='new')],['Q1'])
    evidence_state.project(a)
    assert a['sources'][0]['evidence_spans'][0]['quote']=='new'
    a['evidence']['claims']=evidence_state.merge_claims(a,[],['Q1'])
    evidence_state.project(a)
    assert not a['sources'][0]['evidence_spans'] and a['evidence']['claims'][0]['status']=='unsupported'
    assert a['research']['summary']==a['evidence']['summary']
    assert '适用性待复核：主张1' in a['research']['summary'] and '已有支持：主张2' in a['research']['summary']


@pytest.mark.parametrize('text,replacement', [('Before. Remove this. After.',''),('Remove this.',''),('Before. Remove this. After.','Limited statement.')])
def test_exclude_edits_and_undo(client,monkeypatch,text,replacement):
    a=patch(client,seeded(client),dict(content=text),'write')
    async def result(*args,**kwargs):return dict(decisions=[dict(issue_id='Q1',wording='Do not use',explanation='User excluded',edits=[dict(target='content',original='Remove this.',replacement=replacement)])])
    monkeypatch.setattr(research,'structured',result)
    body=dict(revision=a['revision'],issue_ids=['Q1'],action='exclude',action_id='exclude-1')
    j=client.post('/api/articles/'+a['id']+'/research/issues/actions',headers=H,json=body).json()['job']
    assert wait(client,j)['status']=='completed'
    a=store.get_article(a['id'])
    assert a['content']==text.replace('Remove this.',replacement)
    assert a['research']['issues'][0]['application_state']=='applied'
    response=client.post('/api/articles/'+a['id']+'/research/issues/actions',headers=H,json=dict(revision=a['revision'],issue_ids=['Q1'],action='undo',action_id='undo-1'))
    assert response.status_code==200,response.text
    assert response.json()['article']['content']==text


def test_partial_application_remains_pending(client,monkeypatch):
    a=patch(client,seeded(client),dict(content='Unique. Repeat. Repeat.'),'write')
    async def result(*args,**kwargs):return dict(decisions=[dict(issue_id='Q1',wording='Do not use',explanation='Ambiguous second anchor',edits=[dict(target='content',original='Unique.',replacement=''),dict(target='content',original='Repeat.',replacement='')])])
    monkeypatch.setattr(research,'structured',result)
    j=client.post('/api/articles/'+a['id']+'/research/issues/actions',headers=H,json=dict(revision=a['revision'],issue_ids=['Q1'],action='exclude',action_id='partial')).json()['job']
    assert wait(client,j)['result']['remaining']==1
    a=store.get_article(a['id'])
    assert a['content']==' Repeat. Repeat.' and a['research']['issues'][0]['application_state']=='partial'
    assert a['materials_state']['pending'][0]['id']=='Q1' and a['stages']['sources']=='needs_input'


def test_snapshots_legacy_compression_corruption_and_backup(client):
    from backend import snapshots
    a=new(client)
    a=patch(client,a,dict(content='long synthetic text '*10000),'write')
    a=patch(client,a,dict(content=a['content']+'latest'),'write')
    version=store.versions(a['id'])[0]
    with store.connection() as db:
        raw=db.execute('SELECT data FROM versions WHERE id=?',(version['id'],)).fetchone()[0]
        assert len(raw)<10000
        legacy=snapshots.decode(raw);assert snapshots.decode(store.encode(legacy))==legacy
        corrupted=json.loads(raw);corrupted['sha256']='bad'
        db.execute('UPDATE versions SET data=? WHERE id=?',(json.dumps(corrupted),version['id']))
    with pytest.raises(ValueError,match='历史版本'):store.restore(a['id'],version['id'],a['revision'])
    assert store.get_article(a['id'])['content']==a['content']
    assert (store.DATA/'backups/before-snapshot-v1.sqlite').is_file()


def test_article_and_version_pagination(client):
    for i in range(23):new(client,'Page '+str(i))
    result=client.get('/api/articles?page=2&page_size=20&query=Page').json()
    assert result['total']==23 and len(result['items'])==3
    assert len(client.get('/api/articles').json())==23
    a=client.get('/api/articles').json()[0]
    for i in range(4):a=patch(client,a,dict(content=str(i)),'write')
    result=client.get('/api/articles/'+a['id']+'/versions?page=2&page_size=3').json()
    assert result['total']==4 and len(result['items'])==1


@pytest.mark.parametrize('quality,boundary,reason,expected',[
    ('suitable','','Direct',True),('limited','Only adults','Direct',True),
    ('limited','','Direct',False),('suitable','','',False),('pending','','Direct',False),
    ('insufficient','Association is not causation','Direct',False)])
def test_fixed_evidence_quality_samples(quality,boundary,reason,expected):
    value=dict(span(),quality=quality,boundary=boundary,adoption_reason=reason)
    assert evidence_state.assessed(value)==expected


def test_scoped_empty_reassessment_does_not_keep_old_resolution(client):
    a=seeded(client)
    rows=evidence_state.merge_issues(a,dict(evidence=[],issues=[]),['Q1'])
    assert rows[0]['status']=='open' and rows[1]['status']=='resolved'
    claims=evidence_state.merge_claims(a,[span('C2','S2','Unrelated new text')],['Q1'])
    assert not claims[0]['evidence'] and claims[1]['evidence']==a['evidence']['claims'][1]['evidence']


@pytest.mark.parametrize('failure',['error','cancel'])
def test_batch_keeps_completed_group_and_counts_unfinished(client,monkeypatch,failure):
    from backend import bounded
    a=seeded(client)
    def seed(v):
        v['content']=' '.join(f'Claim-{i}.' for i in range(21))
        v['research']['issues']=[dict(id=f'Q{i}',text=f'Question {i}',kind='blocking',status='open',source_ids=[],claim=f'Claim-{i}') for i in range(21)]
    a=store.save_article(a['id'],a['revision'],seed,'batch fixture')
    ids=[f'Q{i}' for i in range(21)];j=store.create_job(a['id'],dict(stage='bound'));calls=[]
    async def result(*args,**kwargs):
        calls.append(1)
        if len(calls)>1:
            if failure=='cancel':raise asyncio.CancelledError()
            raise ValueError('Synthetic second-group failure')
        return dict(decisions=[dict(issue_id=f'Q{i}',wording='Exclude',explanation='Remove',edits=[dict(target='content',original=f'Claim-{i}.',replacement='')]) for i in range(10)])
    monkeypatch.setattr(research,'structured',result)
    asyncio.run(bounded.run(j['id'],a,ids,'exclude'))
    job=store.job(j['id']);saved=store.get_article(a['id'])
    assert job['status']==('cancelled' if failure=='cancel' else 'failed')
    assert len(calls)==2 and job['result']['remaining']==11 and job['result']['changed']==10
    assert len(saved['research_decisions'])==10 and 'Claim-10.' in saved['content']
    assert 'Claim-0.' not in saved['content']


@pytest.mark.parametrize('related',[False,True])
def test_local_edit_preserves_interleaved_manual_changes(client,monkeypatch,related):
    from backend import bounded
    a=patch(client,seeded(client),dict(content='Start. Target. End.'),'write')
    j=store.create_job(a['id'],dict(stage='bound'))
    async def result(*args,**kwargs):
        current=store.get_article(a['id'])
        store.save_article(a['id'],current['revision'],lambda v:v.update(content='Start. Manually changed. End.' if related else 'Manual prefix. Start. Target. End.'),'manual')
        return dict(decisions=[dict(issue_id='Q1',wording='Exclude',explanation='Remove target',edits=[dict(target='content',original='Target.',replacement='')])])
    monkeypatch.setattr(research,'structured',result)
    asyncio.run(bounded.run(j['id'],a,['Q1'],'exclude'));saved=store.get_article(a['id'])
    assert saved['research_decisions']['Q1']['application_state']==('pending' if related else 'applied')
    if related:assert saved['content']=='Start. Manually changed. End.'
    else:
        restored=store.save_article(a['id'],saved['revision'],lambda v:bounded.undo(v,['Q1']),'undo')
        assert restored['content']=='Manual prefix. Start. Target. End.'


def test_undo_refuses_changed_deletion_context_and_overlaps(client):
    from backend import bounded
    a=patch(client,seeded(client),dict(content='Before. Target. After.'),'write')
    edit=dict(target='content',original='Target.',replacement='')
    assert bounded.replace(a,edit)
    a['content']='Changed context. After.'
    assert not bounded.replace(a,edit,undo=True)
    a['content']='ABC DEF';first=dict(target='content',original='ABC',replacement='X')
    assert bounded.replace(a,first)
    assert not bounded.replace(a,dict(target='content',original='ABC DEF',replacement='Y'))


def test_backup_failure_prevents_snapshot_and_article_write(client,monkeypatch):
    from backend import snapshots
    a=new(client);before=store.get_article(a['id'])
    def fail(*args,**kwargs):raise OSError('Synthetic disk error')
    with monkeypatch.context() as m:
        m.setattr(snapshots.sqlite3,'connect',fail)
        with pytest.raises(ValueError,match='副本'):store.save_article(a['id'],a['revision'],lambda v:v.update(content='must not save'),'change')
    assert store.get_article(a['id'])==before and not store.versions(a['id'])


def test_settings_database_failure_rolls_back_credentials(client):
    value=Settings.model_validate(dict(services=[dict(id='a',model='m',key='old')],default_service='a'))
    providers.save_settings(value);before=store.get_settings()
    with store.connection() as db:
        db.execute("CREATE TRIGGER fail_settings BEFORE INSERT ON settings BEGIN SELECT RAISE(ABORT, 'synthetic failure'); END")
    value.services[0].key='new'
    with pytest.raises(sqlite3.IntegrityError):providers.save_settings(value)
    assert security.key('a')=='old' and store.get_settings()==before


def test_review_invalidates_on_evidence_and_decisions_not_presentation(client):
    from backend import review_state,flow_state
    a=seeded(client);a['stages']['review']='done';a['review']=dict(reviewed_key=review_state.signature(a),completion='ai')
    before=copy.deepcopy(a);flow_state.present(a);flow_state.present(a)
    assert a['stages']['review']=='done'
    a['evidence']['claims'][0]['evidence'][0]['use_scope']='Changed scope'
    flow_state.present(a);assert a['stages']['review']=='stale'
    a=before;a['research_decisions']={'Q1':dict(handling='excluded',wording='Exclude')}
    review_state.present(a);assert a['stages']['review']=='stale'


def test_limits_precedence_for_single_and_batch_verification(client,monkeypatch):
    from backend.models import IssueAction
    from backend import issue_actions
    a=patch(client,seeded(client),dict(research_limits=dict(max_calls=33,max_pages=40,max_rounds=3)))
    requests=[]
    def start(aid,request):requests.append(request);return None
    monkeypatch.setattr(workflow,'start',start)
    for ids in (['Q1'],['Q1','Q2']):
        issue_actions.apply(a['id'],IssueAction(revision=a['revision'],action='verify',issue_ids=ids,action_id=str(ids)))
        job=store.create_job(a['id'],requests[-1].model_dump())
        worker=research.Research(a,job['id'],'research')
        assert worker.cfg['max_calls']==33
        store.update_job(job['id'],status='completed')
    job=store.create_job(a['id'],dict(stage='research',research_limits=dict(max_calls=7,max_pages=8,max_rounds=1)))
    assert research.Research(a,job['id'],'research').cfg['max_calls']==7


def test_old_evidence_missing_assessment_is_only_marked_pending(client):
    a=seeded(client);a['research'].pop('policy_version');a['evidence']['claims'][0]['evidence'][0].pop('adoption_reason')
    evidence_state.sync(a)
    assert a['evidence']['claims'][0]['assessment_pending']
    assert a['research']['issues'][0]['status']=='open' and a['research']['issues'][1]['status']=='resolved'


@pytest.mark.parametrize('value',[[],{},dict(id='fake',outline=[])])
def test_snapshot_invalid_structure_is_rejected(value):
    from backend import snapshots
    for raw in (json.dumps(value),snapshots.encode(value)):
        with pytest.raises(ValueError,match='历史版本'):snapshots.decode(raw)


def test_docx_heading_image_warning_and_image_only_failure():
    from PIL import Image
    blob=io.BytesIO();Image.new('RGB',(20,20),'white').save(blob,'PNG')
    doc=Document();doc.add_heading('Results',level=2);doc.add_picture(io.BytesIO(blob.getvalue()))
    content=io.BytesIO();doc.save(content)
    text,_=materials.extract_file('mixed.docx',content.getvalue())
    assert text.startswith('## Results') and '图片未进行文字或图表识别' in text
    doc=Document();doc.add_picture(io.BytesIO(blob.getvalue()));content=io.BytesIO();doc.save(content)
    with pytest.raises(ValueError,match='扫描'):materials.extract_file('scan.docx',content.getvalue())


def test_scoped_result_cache_keeps_scope_after_save(client,monkeypatch):
    a=seeded(client);job=store.create_job(a['id'],dict(stage='research',issue_ids=['Q1']))
    signatures=[]
    async def run(self,query):
        signatures.append(research.digest([self.cfg,None]))
        self.notes=dict(summary='Scoped',evidence=[span()],issues=[a['research']['issues'][0]],gaps=[],conflicts=[],followup_queries=[])
        return False
    monkeypatch.setattr(research.Research,'run',run)
    saved,_=asyncio.run(research.gather(a,job['id'],'research'))
    assert saved['research']['input_key']==research.input_key(saved,'research','',signatures[0],['Q1'])
    asyncio.run(research.gather(saved,job['id'],'research'))
    assert len(signatures)==1


def test_generic_gap_can_add_first_evidence_without_replacing_unrelated(client):
    a=seeded(client)
    a['research']['issues'].append(dict(id='empty',text='Missing material',claim='',source_ids=[],status='open',kind='blocking'))
    claims=evidence_state.merge_claims(a,[span('C3','S1')],['empty'])
    assert len(claims)==3 and claims[:2]==a['evidence']['claims']
